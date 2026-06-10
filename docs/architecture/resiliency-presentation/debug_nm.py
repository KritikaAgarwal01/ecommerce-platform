from pptx import Presentation
from lxml import etree

pptx_path = r"c:/Users/ambarana/ecommerce-platform/docs/architecture/resiliency-presentation/slide-deck/resiliency-research-skill.pptx"

prs = Presentation(pptx_path)

# Check notes master placeholders
notes_master_part = prs.part.notes_master_part
nm = notes_master_part.notes_master
print("Notes master placeholders:")
for ph in nm.placeholders:
    print(f"  idx={ph.placeholder_format.idx} type={ph.placeholder_format.type} ph_type={ph.element.ph_type}")

print("\nNotes master XML (first 3000 chars):")
nm_xml = etree.tostring(nm._element, pretty_print=True).decode()
print(nm_xml[:3000])

from pptx import Presentation
import copy

pptx_path = r"c:/Users/ambarana/ecommerce-platform/docs/architecture/resiliency-presentation/slide-deck/resiliency-research-skill.pptx"
out_path  = r"c:/Users/ambarana/ecommerce-platform/docs/architecture/resiliency-presentation/slide-deck/resiliency-research-skill-prep.pptx"

prs = Presentation(pptx_path)
print(f"Current slides: {len(prs.slides)}")

# Capture placeholder elements from slide 1's working notes slide as a template.
# PP_PLACEHOLDER.BODY==2, so python-pptx finds the idx=2 (sldImg) sp as the 'body'.
# We need ALL three placeholder sp elements to be present on new slides.
template_sps = [
    el for el in prs.slides[0].notes_slide.shapes._spTree.iterchildren()
    if el.tag.endswith("}sp")
]
print(f"Template notes sp count: {len(template_sps)}")

# Trim to exactly 12 slides
while len(prs.slides) > 12:
    rId = prs.slides._sldIdLst[-1].rId
    prs.part.drop_rel(rId)
    prs.slides._sldIdLst.remove(prs.slides._sldIdLst[-1])
print(f"After trim: {len(prs.slides)}")

# Add 2 blank slides (positions 13 and 14) with cloned notes placeholders
blank_layout = prs.slide_layouts[6]
for i in range(2):
    s = prs.slides.add_slide(blank_layout)
    ns = s.notes_slide
    sp_tree = ns.shapes._spTree
    for tmpl_sp in template_sps:
        new_sp = copy.deepcopy(tmpl_sp)
        # Clear text runs so notes body is empty; build_deck will populate it
        A = "http://schemas.openxmlformats.org/drawingml/2006/main"
        for r_el in new_sp.findall(f".//{{{A}}}r"):
            t_el = r_el.find(f"{{{A}}}t")
            if t_el is not None:
                t_el.text = ""
        sp_tree.append(new_sp)
    ntf = ns.notes_text_frame
    print(f"Slide {13+i}: notes_text_frame = {type(ntf).__name__}")
    if ntf is not None:
        print(f"  text preview: '{ntf.text[:40]}'")

print(f"Final slides: {len(prs.slides)}")
prs.save(out_path)
print(f"Saved to {out_path}")

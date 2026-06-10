from pptx import Presentation
from lxml import etree

pptx_path = r"c:/Users/ambarana/ecommerce-platform/docs/architecture/resiliency-presentation/slide-deck/resiliency-research-skill.pptx"

prs = Presentation(pptx_path)
print(f"Current slides: {len(prs.slides)}")

# Inspect notes slide for slide 1 (should be working)
s1 = prs.slides[0]
ns1 = s1.notes_slide
print(f"\nSlide 1 notes_slide type: {type(ns1).__name__}")
print(f"Slide 1 notes_text_frame: {ns1.notes_text_frame}")
print(f"Slide 1 notes placeholders: {list(ns1.placeholders)}")
print(f"Slide 1 notes shapes count: {len(ns1.shapes)}")

# Print the notes slide XML
print("\nSlide 1 notes XML (first 2000 chars):")
ns1_xml = etree.tostring(ns1._element, pretty_print=True).decode()
print(ns1_xml[:2000])

# Now add a blank slide and inspect
blank_layout = prs.slide_layouts[6]
snew = prs.slides.add_slide(blank_layout)
nsnew = snew.notes_slide
print(f"\nNew slide notes_slide type: {type(nsnew).__name__}")
print(f"New slide notes_text_frame: {nsnew.notes_text_frame}")
print(f"New slide notes placeholders: {list(nsnew.placeholders)}")
print(f"New slide notes shapes count: {len(nsnew.shapes)}")
print("\nNew slide notes XML (first 2000 chars):")
nsnew_xml = etree.tostring(nsnew._element, pretty_print=True).decode()
print(nsnew_xml[:2000])
